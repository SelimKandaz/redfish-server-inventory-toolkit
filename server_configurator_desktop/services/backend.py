from __future__ import annotations

import os
import re
from collections import Counter, defaultdict
from io import BytesIO
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

import requests
import urllib3
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
try:
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import BooleanObject, NameObject
except ImportError:
    from PyPDF2 import PdfReader, PdfWriter
    from PyPDF2.generic import BooleanObject, NameObject
from reportlab.lib import colors
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from ..models import TaskRow

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

LogFn = Callable[[str], None]
ProgressFn = Callable[[int, int], None]
StatusFn = Callable[[str, str], None]


@dataclass
class RunOptions:
    timeout_seconds: int = 30
    skip_unreachable: bool = True
    create_individual_pdfs: bool = True
    create_group_pdf: bool = True
    create_serial_excel: bool = True


class BackendError(RuntimeError):
    pass


class ConfiguratorBackend:
    FIELD_MAP = {
        "CHASSIS_PLATFORM": "Text4",
        "QUANTITY": "Text5",
        "MOTHERBOARD": "Text6",
        "SYSTEM_SERIALS": "Text8",
        "CPU": "Text9",
        "MEMORY": "Text10",
        "DRIVES": "Text12",
        "RAID_HW": "Text16",
        "NETWORK_HW": "Text17",
        "PSU": "Text19",
        "EXTRAS": "Text20",
        "BIOS": "Text21",
        "BMC_ILO": "Text22",
        "NOTES": None,
    }

    def __init__(self, base_dir: Path, template_pdf: Path):
        self.base_dir = Path(base_dir)
        self.template_pdf = Path(template_pdf)


    def _normalized(self, value) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def _is_placeholder(self, value) -> bool:
        text = self._normalized(value).upper()
        return text in {"", "NOT AVAILABLE", "N/A", "NA", "NONE", "UNKNOWN", "NULL"}

    def _clean_part_number(self, value) -> str:
        text = self._normalized(value)
        return "" if self._is_placeholder(text) else text

    def _clean_serial(self, value) -> str:
        text = self._normalized(value)
        return "" if self._is_placeholder(text) else text

    def _pretty_counter(self, counter: Counter) -> str:
        if not counter:
            return "none"
        return "; ".join(f"{qty}x {key}" for key, qty in sorted(counter.items(), key=lambda item: str(item[0])))

    def rf_get(self, ip: str, path: str, auth: tuple[str, str], timeout: int, log: LogFn) -> dict:
        url = path if path.startswith("http") else f"https://{ip}{path}"
        try:
            response = requests.get(url, auth=auth, verify=False, timeout=timeout)
            if response.status_code == 404:
                return {}
            response.raise_for_status()
            return response.json()
        except requests.exceptions.ConnectTimeout:
            log(f"[{ip}] TIMEOUT {path}")
            return {}
        except requests.exceptions.ReadTimeout:
            log(f"[{ip}] READ TIMEOUT {path}")
            return {}
        except requests.exceptions.HTTPError as exc:
            log(f"[{ip}] GET {path} failed: {exc}")
            return {}
        except requests.exceptions.RequestException as exc:
            log(f"[{ip}] GET {path} failed: {exc}")
            return {}
        except ValueError as exc:
            log(f"[{ip}] Invalid JSON at {path}: {exc}")
            return {}

    def rf_first_member_path(self, collection: dict) -> str | None:
        members = collection.get("Members") or []
        if not members:
            return None
        return members[0].get("@odata.id")

    def discover_paths(self, ip: str, auth: tuple[str, str], vendor_key: str, timeout: int, log: LogFn) -> tuple[str, str, str]:
        sys_coll = self.rf_get(ip, "/redfish/v1/Systems", auth, timeout, log)
        sys_path = self.rf_first_member_path(sys_coll)
        if not sys_path:
            sys_path = "/redfish/v1/Systems/System.Embedded.1" if vendor_key == "dell" else "/redfish/v1/Systems/1"

        ch_coll = self.rf_get(ip, "/redfish/v1/Chassis", auth, timeout, log)
        ch_path = self.rf_first_member_path(ch_coll)
        if not ch_path:
            ch_path = "/redfish/v1/Chassis/System.Embedded.1" if vendor_key == "dell" else "/redfish/v1/Chassis/1"

        mgr_coll = self.rf_get(ip, "/redfish/v1/Managers", auth, timeout, log)
        mgr_path = self.rf_first_member_path(mgr_coll) or "/redfish/v1/Managers/1"
        return sys_path, ch_path, mgr_path

    def get_system(self, ip, auth, sys_path, timeout, log):
        return self.rf_get(ip, sys_path, auth, timeout, log)

    def get_chassis(self, ip, auth, ch_path, timeout, log):
        return self.rf_get(ip, ch_path, auth, timeout, log)

    def get_bios_version(self, system: dict) -> str:
        for key in ("BiosVersion", "BIOSVersion", "Bios"):
            value = system.get(key)
            if isinstance(value, str) and value:
                return value
            if isinstance(value, dict):
                for inner in ("Current", "Version", "VersionString"):
                    if isinstance(value.get(inner), str) and value.get(inner):
                        return value.get(inner)
        oem = system.get("Oem") or {}
        for section in oem.values():
            if isinstance(section, dict):
                for name, value in section.items():
                    if isinstance(value, str) and "bios" in name.lower():
                        return value
        return ""

    def get_bmc_version(self, ip, auth, mgr_path, timeout, log) -> str:
        mgr = self.rf_get(ip, mgr_path, auth, timeout, log)
        if not mgr:
            return ""
        for key in ("FirmwareVersion", "ManagerFirmwareVersion", "Version"):
            value = mgr.get(key)
            if isinstance(value, str) and value:
                return value
        oem = mgr.get("Oem") or {}
        for section in oem.values():
            if isinstance(section, dict):
                for name, value in section.items():
                    if isinstance(value, str) and ("fw" in name.lower() or "firmware" in name.lower()):
                        return value
        return ""

    def get_processors(self, ip, auth, sys_path, timeout, log):
        out = []
        coll = self.rf_get(ip, f"{sys_path}/Processors", auth, timeout, log)
        for member in coll.get("Members") or []:
            proc = self.rf_get(ip, member.get("@odata.id"), auth, timeout, log)
            if proc:
                out.append(proc)
        return out

    def get_dimms(self, ip, auth, sys_path, timeout, log):
        out = []
        coll = self.rf_get(ip, f"{sys_path}/Memory", auth, timeout, log)
        for member in coll.get("Members") or []:
            dimm = self.rf_get(ip, member.get("@odata.id"), auth, timeout, log)
            if dimm:
                out.append(dimm)
        return out

    def get_drives(self, ip, auth, sys_path, ch_path, timeout, log):
        out = []

        def add_drive_like(drive: dict):
            if not drive:
                return
            serial = drive.get("SerialNumber")
            if not serial:
                for ident in drive.get("Identifiers") or []:
                    if isinstance(ident, dict):
                        fmt = (ident.get("DurableNameFormat") or "").lower()
                        if "serial" in fmt or "sn" in fmt:
                            serial = ident.get("DurableName")
                            if serial:
                                break
            if not serial:
                oem = drive.get("Oem") or {}
                for section in oem.values():
                    if isinstance(section, dict):
                        for name, value in section.items():
                            if isinstance(value, str) and "serial" in name.lower():
                                serial = value
                                break
                    if serial:
                        break
            drive = dict(drive)
            drive["_BetterSerial"] = serial
            out.append(drive)

        store_coll = self.rf_get(ip, f"{sys_path}/Storage", auth, timeout, log)
        for member in store_coll.get("Members") or []:
            storage = self.rf_get(ip, member.get("@odata.id"), auth, timeout, log)
            for ref in storage.get("Drives") or []:
                drive = self.rf_get(ip, ref.get("@odata.id"), auth, timeout, log)
                add_drive_like(drive)

        if ch_path:
            chassis = self.rf_get(ip, ch_path, auth, timeout, log)
            for ref in chassis.get("Drives") or []:
                drive = self.rf_get(ip, ref.get("@odata.id"), auth, timeout, log)
                add_drive_like(drive)

            ch_drives = self.rf_get(ip, f"{ch_path}/Drives", auth, timeout, log)
            for member in ch_drives.get("Members") or []:
                drive = self.rf_get(ip, member.get("@odata.id"), auth, timeout, log)
                add_drive_like(drive)

            pcie = self.rf_get(ip, f"{ch_path}/PCIeDevices", auth, timeout, log)
            for member in pcie.get("Members") or []:
                device = self.rf_get(ip, member.get("@odata.id"), auth, timeout, log)
                if not device:
                    continue
                dev_type = (device.get("DeviceType") or "").lower()
                model = (device.get("Model") or device.get("Name") or "").lower()
                if not any(token in dev_type for token in ["nvme", "ssd", "storage"]):
                    if "nvme" not in model and "ssd" not in model:
                        continue
                add_drive_like({
                    "Model": device.get("Model") or device.get("Name"),
                    "PartNumber": device.get("PartNumber"),
                    "SerialNumber": device.get("SerialNumber"),
                    "CapacityBytes": device.get("CapacityBytes"),
                    "Oem": device.get("Oem"),
                    "Identifiers": device.get("Identifiers"),
                })
        return out

    def get_nic_ports(self, ip, auth, sys_path, mgr_path, timeout, log):
        out = []
        coll = self.rf_get(ip, f"{sys_path}/EthernetInterfaces", auth, timeout, log)
        for member in coll.get("Members") or []:
            nic = self.rf_get(ip, member.get("@odata.id"), auth, timeout, log)
            if nic:
                out.append(nic)
        if out:
            return out
        coll2 = self.rf_get(ip, f"{mgr_path}/EthernetInterfaces", auth, timeout, log)
        for member in coll2.get("Members") or []:
            nic = self.rf_get(ip, member.get("@odata.id"), auth, timeout, log)
            if nic:
                out.append(nic)
        return out

    def get_manager_ports(self, ip, auth, mgr_path, timeout, log):
        out = []
        coll = self.rf_get(ip, f"{mgr_path}/EthernetInterfaces", auth, timeout, log)
        for member in coll.get("Members") or []:
            nic = self.rf_get(ip, member.get("@odata.id"), auth, timeout, log)
            if nic:
                out.append(nic)
        return out

    def get_nic_adapters(self, ip, auth, sys_path, ch_path, timeout, log):
        out = []
        coll = self.rf_get(ip, f"{ch_path}/NetworkAdapters", auth, timeout, log)
        for member in coll.get("Members") or []:
            nic = self.rf_get(ip, member.get("@odata.id"), auth, timeout, log)
            if nic:
                out.append(nic)
        coll2 = self.rf_get(ip, f"{sys_path}/NetworkAdapters", auth, timeout, log)
        for member in coll2.get("Members") or []:
            nic = self.rf_get(ip, member.get("@odata.id"), auth, timeout, log)
            if nic:
                out.append(nic)
        return out

    def get_psus(self, ip, auth, ch_path, timeout, log):
        out = []
        power = self.rf_get(ip, f"{ch_path}/Power", auth, timeout, log)
        for psu in power.get("PowerSupplies") or []:
            if not isinstance(psu, dict):
                continue
            serial = psu.get("SerialNumber")
            if not serial:
                oem = psu.get("Oem") or {}
                for section in oem.values():
                    if isinstance(section, dict):
                        for name, value in section.items():
                            if isinstance(value, str) and "serial" in name.lower():
                                serial = value
                                break
                    if serial:
                        break
            out.append({
                "SerialNumber": serial,
                "Model": psu.get("Model"),
                "CapacityWatts": psu.get("CapacityWatts") or psu.get("PowerCapacityWatts"),
                "LineInputVoltageType": psu.get("LineInputVoltageType"),
            })
        return out

    def summarize_cpu(self, procs):
        if not procs:
            return ""
        groups = Counter()
        for proc in procs:
            key = (proc.get("Model") or "", proc.get("PartNumber") or "")
            groups[key] += 1
        parts = []
        for (model, pn), count in groups.items():
            if not model and not pn:
                continue
            txt = f"{count}x " if count > 1 else ""
            txt += model
            if pn:
                txt += f" ({pn})" if model else pn
            parts.append(txt)
        return "; ".join(parts)

    def summarize_memory(self, dimms):
        if not dimms:
            return ""
        total_mib = 0
        by_pn = Counter()
        sample_size = {}
        for dimm in dimms:
            cap = dimm.get("CapacityMiB") or 0
            pn = self._clean_part_number(dimm.get("PartNumber") or dimm.get("Manufacturer") or "")
            if not cap or not pn:
                continue
            total_mib += cap
            by_pn[pn] += 1
            if pn not in sample_size:
                sample_size[pn] = int(cap) // 1024
        if not by_pn:
            return ""
        total_gb = round(total_mib / 1024.0, 1)
        parts = [f"Total: {total_gb}GB"]
        for pn, qty in by_pn.items():
            size = sample_size.get(pn)
            parts.append(f"{qty}x {size}GB ({pn})" if size else f"{qty}x ({pn})")
        return "; ".join(parts)

    def summarize_psus(self, psus):
        if not psus:
            return ""
        groups = Counter()
        for psu in psus:
            watts = psu.get("CapacityWatts")
            model = self._clean_part_number(psu.get("Model"))
            if not watts and not model:
                continue
            key = (watts, psu.get("LineInputVoltageType"), model)
            groups[key] += 1
        parts = []
        for (watts, volt, model), count in groups.items():
            txt = f"{count}x " if count > 1 else ""
            if watts:
                txt += f"{watts}W"
            if volt:
                txt += f" {volt}"
            if model:
                txt += f" ({model})" if txt else model
            if txt.strip():
                parts.append(txt.strip())
        return "; ".join(parts)

    def summarize_drives(self, drives):
        if not drives:
            return ""
        groups = Counter()
        for drive in drives:
            model = self._normalized(drive.get("Model") or "Drive")
            pn = self._clean_part_number(drive.get("PartNumber") or "")
            size = drive.get("CapacityBytes")
            size_gb = None
            if isinstance(size, int) and size > 0:
                size_gb = int(round(size / (1024 ** 3)))
            groups[(pn, model, size_gb)] += 1
        parts = []
        for (pn, model, size_gb), qty in groups.items():
            base = f"{qty}x {size_gb}GB {model}" if size_gb else f"{qty}x {model}"
            if pn:
                base += f" ({pn})"
            parts.append(base)
        return "; ".join(parts)

    def summarize_network(self, nic_adapters):
        if not nic_adapters:
            return ""
        groups = Counter()
        for adapter in nic_adapters:
            model = self._normalized(adapter.get("Model") or adapter.get("Name") or "NIC")
            pn = self._clean_part_number(adapter.get("PartNumber") or "")
            if not pn:
                oem = adapter.get("Oem") or {}
                for section in oem.values():
                    if isinstance(section, dict):
                        for name, value in section.items():
                            if isinstance(value, str) and ("part" in name.lower() or "spn" in name.lower()):
                                pn = self._clean_part_number(value)
                                break
                    if pn:
                        break
            groups[(model, pn)] += 1
        parts = []
        for (model, pn), qty in groups.items():
            txt = f"{qty}x " if qty > 1 else ""
            txt += model
            if pn:
                txt += f" ({pn})"
            parts.append(txt)
        return "; ".join(parts)

    def describe_for_pdf(self, server: dict) -> dict:
        system = server.get("system") or {}
        chassis = server.get("chassis") or {}
        return {
            "CHASSIS_PLATFORM": chassis.get("Model") or system.get("Model") or "",
            "MOTHERBOARD": system.get("PartNumber") or chassis.get("PartNumber") or chassis.get("Model") or system.get("Model") or "",
            "SYSTEM_SERIALS": system.get("SerialNumber") or chassis.get("SerialNumber") or "",
            "CPU": self.summarize_cpu(server.get("procs") or []),
            "MEMORY": self.summarize_memory(server.get("dimms") or []),
            "DRIVES": self.summarize_drives(server.get("drives") or []),
            "RAID_HW": "",
            "NETWORK_HW": self.summarize_network(server.get("nic_adapters") or []),
            "PSU": self.summarize_psus(server.get("psus") or []),
            "BIOS": (server.get("bios") or "").strip(),
            "BMC_ILO": (server.get("bmcver") or "").strip(),
            "EXTRAS": "",
            "QUANTITY": "1",
        }

    def _field_rects(self, reader: PdfReader) -> dict[str, list[float]]:
        page = reader.pages[0]
        rects: dict[str, list[float]] = {}
        annots = page.get("/Annots")
        if annots is None:
            return rects
        annots = annots.get_object()
        for annot_ref in annots:
            annot = annot_ref.get_object()
            if annot.get("/Subtype") != "/Widget":
                continue
            name = annot.get("/T")
            rect = annot.get("/Rect")
            if name and rect:
                rects[str(name)] = [float(v) for v in rect]
            elif rect:
                vals = [float(v) for v in rect]
                x0, y0, x1, y1 = vals
                if x0 > 180 and y0 < 70 and y1 < 120:
                    rects["__NOTES__"] = vals
        return rects

    def _draw_text_to_rect(self, c, text: str, rect: list[float], *, font_name: str = "Helvetica", font_size: float = 12, bold: bool = False, multiline: bool = False, align: str = "left"):
        if not text:
            return
        x0, y0, x1, y1 = rect
        pad_x = 10
        pad_y = 6
        width = max(10, x1 - x0 - pad_x * 2)
        height = max(10, y1 - y0 - pad_y * 2)
        if bold:
            font_name = "Helvetica-Bold"
        if multiline:
            size = font_size
            lines = simpleSplit(text, font_name, size, width)
            while size > 7 and len(lines) * (size + 2) > height:
                size -= 1
                lines = simpleSplit(text, font_name, size, width)
            text_obj = c.beginText()
            text_obj.setFont(font_name, size)
            text_obj.setLeading(size + 2)
            text_obj.setTextOrigin(x0 + pad_x, y1 - pad_y - size)
            for line in lines:
                text_obj.textLine(line)
            c.drawText(text_obj)
            return

        size = font_size
        while size > 7 and stringWidth(text, font_name, size) > width:
            size -= 0.5
        y = y0 + (height - size) / 2 + pad_y
        if align == "center":
            tx = x0 + (x1 - x0 - stringWidth(text, font_name, size)) / 2
        else:
            tx = x0 + pad_x
        c.setFont(font_name, size)
        c.drawString(tx, y, text)

    def write_pdf_from_values(self, values: dict, output_pdf: Path):
        if not self.template_pdf.exists():
            raise BackendError(f"Template PDF not found: {self.template_pdf}")

        reader = PdfReader(str(self.template_pdf))
        rects = self._field_rects(reader)
        page0 = reader.pages[0]
        page_w = float(page0.mediabox.width)
        page_h = float(page0.mediabox.height)

        overlay_stream = BytesIO()
        c = canvas.Canvas(overlay_stream, pagesize=(page_w, page_h))

        field_text = {
            self.FIELD_MAP.get("CHASSIS_PLATFORM"): values.get("CHASSIS_PLATFORM", ""),
            self.FIELD_MAP.get("QUANTITY"): values.get("QUANTITY", ""),
            self.FIELD_MAP.get("MOTHERBOARD"): values.get("MOTHERBOARD", ""),
            self.FIELD_MAP.get("SYSTEM_SERIALS"): values.get("SYSTEM_SERIALS", ""),
            self.FIELD_MAP.get("CPU"): values.get("CPU", ""),
            self.FIELD_MAP.get("MEMORY"): values.get("MEMORY", ""),
            self.FIELD_MAP.get("DRIVES"): values.get("DRIVES", ""),
            self.FIELD_MAP.get("RAID_HW"): values.get("RAID_HW", ""),
            self.FIELD_MAP.get("NETWORK_HW"): values.get("NETWORK_HW", ""),
            self.FIELD_MAP.get("PSU"): values.get("PSU", ""),
            self.FIELD_MAP.get("EXTRAS"): values.get("EXTRAS", ""),
            self.FIELD_MAP.get("BIOS"): values.get("BIOS", ""),
            self.FIELD_MAP.get("BMC_ILO"): values.get("BMC_ILO", ""),
            "__NOTES__": values.get("NOTES", ""),
        }

        single_line_sizes = {
            "Text4": 10.5, "Text5": 10.5, "Text6": 10.0, "Text8": 10.0, "Text9": 9.0,
            "Text10": 9.0, "Text14": 9.0, "Text15": 9.0, "Text16": 9.0, "Text17": 9.0,
            "Text18": 9.0, "Text19": 9.0, "Text20": 9.0, "Text21": 9.0, "Text22": 9.0,
        }
        multiline_fields = {"Text12", "Text20", "__NOTES__"}

        for field_name, value in field_text.items():
            if not field_name or not value or field_name not in rects:
                continue
            self._draw_text_to_rect(
                c,
                str(value),
                rects[field_name],
                font_size=single_line_sizes.get(field_name, 9.0),
                multiline=(field_name in multiline_fields),
            )

        c.save()
        overlay_stream.seek(0)
        overlay_reader = PdfReader(overlay_stream)
        overlay_page = overlay_reader.pages[0]

        writer = PdfWriter()
        base_page = reader.pages[0]
        base_page.merge_page(overlay_page)
        # flatten by removing interactive form usage from output
        writer.add_page(base_page)

        output_pdf.parent.mkdir(parents=True, exist_ok=True)
        with output_pdf.open("wb") as handle:
            writer.write(handle)



    def _finalize_fillable_pdf(self, pdf_path: Path):
        try:
            reader = PdfReader(str(pdf_path))
            writer = PdfWriter()
            for page in reader.pages:
                writer.add_page(page)
            try:
                if "/AcroForm" in reader.trailer["/Root"]:
                    writer._root_object.update({NameObject("/AcroForm"): reader.trailer["/Root"]["/AcroForm"]})
                    writer._root_object["/AcroForm"].update({NameObject("/NeedAppearances"): BooleanObject(True)})
            except Exception:
                pass
            with pdf_path.open("wb") as handle:
                writer.write(handle)
        except Exception:
            pass

    def write_pdf_from_values_editable(self, values: dict, output_pdf: Path):
        """
        Create a fillable review PDF with visible default text for manual cleanup.
        """
        output_pdf.parent.mkdir(parents=True, exist_ok=True)

        page_width = 792
        page_height = 612
        try:
            reader = PdfReader(str(self.template_pdf))
            page = reader.pages[0]
            page_width = float(page.mediabox.width)
            page_height = float(page.mediabox.height)
        except Exception:
            pass

        c = canvas.Canvas(str(output_pdf), pagesize=(page_width, page_height))
        form = c.acroForm

        def label(x, y, text, size=10):
            c.setFont("Helvetica-Bold", size)
            c.drawString(x, y, text)

        def field(name, x, y, w, h, value="", multiline=False, size=10):
            flags = 4096 if multiline else 0
            # draw current text so the PDF never looks blank
            c.setFont("Helvetica", size)
            text_value = str(value or "")
            if multiline:
                lines = simpleSplit(text_value, "Helvetica", size, max(w - 8, 20))
                baseline = y + h - size - 4
                for line in lines[: max(int((h - 8) / (size + 2)), 1)]:
                    c.drawString(x + 4, baseline, line)
                    baseline -= size + 2
            else:
                c.drawString(x + 4, y + max((h - size) / 2, 1), text_value)

            form.textfield(
                name=name,
                x=x, y=y, width=w, height=h,
                value=text_value,
                fontName="Helvetica",
                fontSize=size,
                borderStyle="solid",
                borderWidth=0.6,
                forceBorder=True,
                textColor=colors.black,
                borderColor=colors.HexColor("#777777"),
                fillColor=None,
                fieldFlags=flags,
            )

        margin = 28
        content_w = page_width - margin * 2
        y = page_height - 34

        c.setFont("Helvetica-Bold", 15)
        c.drawString(margin, y, "Configurator Group Review")
        y -= 18
        c.setFont("Helvetica", 9)
        c.drawString(margin, y, "Editable review copy. Remove or adjust lines as needed before sharing.")
        y -= 22

        blocks = [
            ("Overview", values.get("EXTRAS", ""), 92),
            ("Configuration Details", values.get("NOTES", ""), 265),
            ("Hardware Summary", "\n".join([
                f"Platform: {values.get('CHASSIS_PLATFORM', '')}",
                f"Quantity: {values.get('QUANTITY', '')}",
                f"Motherboard: {values.get('MOTHERBOARD', '')}",
                f"System Serials: {values.get('SYSTEM_SERIALS', '')}",
                f"CPU: {values.get('CPU', '')}",
                f"Memory: {values.get('MEMORY', '')}",
                f"Drives: {values.get('DRIVES', '')}",
                f"Network: {values.get('NETWORK_HW', '')}",
                f"PSU: {values.get('PSU', '')}",
                f"BIOS: {values.get('BIOS', '')}",
                f"BMC/iLO: {values.get('BMC_ILO', '')}",
            ]), 150),
        ]

        for idx, (caption, block_text, height) in enumerate(blocks, start=1):
            label(margin, y, caption, 10)
            field(f"block_{idx}", margin, y - height, content_w, height - 12, block_text, multiline=True, size=9)
            y -= height + 24
            if y < 90 and idx != len(blocks):
                c.showPage()
                form = c.acroForm
                y = page_height - 34

        c.showPage()
        c.save()
        self._finalize_fillable_pdf(output_pdf)

    def build_serial_columns(self, servers: list[dict]) -> dict[str, list[str]]:
        headers = [
            "Chassis SN:", "MB SN:", "Blade SN:", "CPU SN:", "Memory SN:",
            "NIC card SN:", "NIC MAC:", "Storage SN:", "PSU SN:", "Server SN:",
        ]
        cols = {h: [] for h in headers}
        for server in servers:
            vendor = (server.get("vendor") or "").lower()
            system = server.get("system") or {}
            chassis = server.get("chassis") or {}
            procs = server.get("procs") or []
            dimms = server.get("dimms") or []
            drives = server.get("drives") or []
            nic_ports = server.get("nic_ports") or []
            nic_adapters = server.get("nic_adapters") or []
            psus = server.get("psus") or []

            sys_sn = system.get("SerialNumber") or ""
            mb_sn = chassis.get("SerialNumber") or ""
            ch_sn = sys_sn if "supermicro" in vendor else ""

            if ch_sn:
                cols["Chassis SN:"].append(ch_sn)
            if mb_sn:
                cols["MB SN:"].append(mb_sn)
            if sys_sn:
                cols["Blade SN:"].append(sys_sn)
                cols["Server SN:"].append(sys_sn)

            for proc in procs:
                serial = proc.get("SerialNumber")
                if serial:
                    cols["CPU SN:"].append(serial)
            for dimm in dimms:
                serial = dimm.get("SerialNumber")
                if serial:
                    cols["Memory SN:"].append(serial)
            seen_nic = set()
            for adapter in nic_adapters:
                serial = adapter.get("SerialNumber")
                if serial and serial not in seen_nic:
                    seen_nic.add(serial)
                    cols["NIC card SN:"].append(serial)
            seen_mac = set()
            for nic in nic_ports:
                mac = nic.get("MACAddress")
                if mac and mac not in seen_mac:
                    seen_mac.add(mac)
                    cols["NIC MAC:"].append(mac)
            for drive in drives:
                serial = drive.get("_BetterSerial") or drive.get("SerialNumber")
                if serial:
                    cols["Storage SN:"].append(serial)
            for psu in psus:
                serial = psu.get("SerialNumber") or psu.get("Model")
                if serial:
                    cols["PSU SN:"].append(serial)
        return cols

    def _style_sheet(self, ws, header_color="1F4E78", freeze="A2", table_name: str | None = None):
        header_fill = PatternFill("solid", fgColor=header_color)
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="D9D9D9")
        alt_fill = PatternFill("solid", fgColor="F7FBFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if row_idx % 2 == 0:
                    cell.fill = alt_fill

        if freeze:
            ws.freeze_panes = freeze

        for column in ws.columns:
            max_len = 0
            col_letter = column[0].column_letter
            for cell in column:
                try:
                    max_len = max(max_len, len(str(cell.value or "")))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)

        if table_name and ws.max_row >= 2 and ws.max_column >= 2:
            seen_headers: dict[str, int] = {}
            for col_idx in range(1, ws.max_column + 1):
                raw = ws.cell(row=1, column=col_idx).value
                header = str(raw).strip() if raw is not None else ""
                if not header:
                    header = f"Column{col_idx}"

                if header in seen_headers:
                    seen_headers[header] += 1
                    header = f"{header}_{seen_headers[header]}"
                else:
                    seen_headers[header] = 1

                ws.cell(row=1, column=col_idx).value = header

            ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            table = Table(displayName=table_name, ref=ref)
            style = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
            table.tableStyleInfo = style
            try:
                ws.add_table(table)
            except Exception:
                ws.auto_filter.ref = ref
        elif ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    def _server_sample_headers(self) -> list[str]:
        return [
            "System Serial Number", "Manufacturer", "Model", "System MAC Address",
            "IPMI IP", "IPMI Username", "IPMI Password",
            "NIC1 MAC1", "NIC1 MAC2", "NIC1 MAC3", "NIC1 MAC4",
            "NIC2 MAC1", "NIC2 MAC2", "NIC2 MAC3", "NIC2 MAC4",
        ]

    def _dedupe_mac_list(self, values: list[str]) -> list[str]:
        out: list[str] = []
        seen: set[str] = set()
        for value in values:
            mac = self._normalized(value).upper()
            if not mac or mac in seen:
                continue
            seen.add(mac)
            out.append(mac)
        return out

    def _guess_nic_group(self, port: dict) -> int | None:
        candidates = [
            port.get("Name"),
            port.get("Id"),
            port.get("Description"),
            port.get("PhysicalPortNumber"),
        ]
        for raw in candidates:
            text = self._normalized(raw)
            if not text:
                continue
            for pattern in [
                r"\bNIC\s*([1-9]\d*)\b",
                r"\bLOM\s*([1-9]\d*)\b",
                r"NIC(?:\.|_|-| )?(?:Integrated|Slot)?(?:\.|_|-| )?([1-9]\d*)",
                r"Port\s*([1-9]\d*)[-/]([1-9]\d*)",
            ]:
                match = re.search(pattern, text, re.IGNORECASE)
                if not match:
                    continue
                if len(match.groups()) >= 2 and pattern.startswith(r"Port"):
                    return int(match.group(1))
                return int(match.group(1))
        links = port.get("Links") or {}
        for raw in (links.get("NetworkAdapter"), links.get("RelatedItem"), links.get("RelatedItems")):
            items = raw if isinstance(raw, list) else [raw]
            for item in items:
                if not isinstance(item, dict):
                    continue
                ref = self._normalized(item.get("@odata.id"))
                if not ref:
                    continue
                match = re.search(r"(?:NetworkAdapters|NIC)(?:/|\.|_|-)?([1-9]\d*)", ref, re.IGNORECASE)
                if match:
                    return int(match.group(1))
        return None

    def _build_sample_server_row(self, server: dict) -> dict:
        system = server.get("system") or {}
        chassis = server.get("chassis") or {}
        manager_ports = self._dedupe_mac_list([
            (port or {}).get("MACAddress") or ""
            for port in (server.get("manager_ports") or [])
        ])
        nic_ports = server.get("nic_ports") or []
        manager_mac_set = {mac.upper() for mac in manager_ports}
        system_port_macs = []
        for port in nic_ports:
            mac = self._normalized((port or {}).get("MACAddress")).upper()
            if not mac:
                continue
            if mac in manager_mac_set:
                continue
            system_port_macs.append(mac)
        system_port_macs = self._dedupe_mac_list(system_port_macs)
        if not system_port_macs:
            system_port_macs = self._dedupe_mac_list([
                (port or {}).get("MACAddress") or ""
                for port in nic_ports
            ])

        grouped: dict[int, list[str]] = {}
        fallback_order: list[str] = []
        for port in nic_ports:
            mac = self._normalized((port or {}).get("MACAddress")).upper()
            if not mac:
                continue
            if manager_mac_set and mac in manager_mac_set:
                continue
            group_no = self._guess_nic_group(port)
            if group_no in (1, 2):
                grouped.setdefault(group_no, [])
                if mac not in grouped[group_no]:
                    grouped[group_no].append(mac)
            else:
                fallback_order.append(mac)

        fallback_order = self._dedupe_mac_list(fallback_order)
        nic1 = list(grouped.get(1, []))
        nic2 = list(grouped.get(2, []))
        if not nic1 and fallback_order:
            nic1 = fallback_order[:4]
            fallback_order = fallback_order[4:]
        if not nic2 and fallback_order:
            nic2 = fallback_order[:4]
        if len(nic1) > 4:
            spill = nic1[4:]
            nic1 = nic1[:4]
            nic2 = (nic2 + spill)[:4]
        if len(nic2) > 4:
            nic2 = nic2[:4]

        def padded(values: list[str]) -> list[str]:
            vals = values[:4]
            while len(vals) < 4:
                vals.append("")
            return vals

        nic1 = padded(nic1)
        nic2 = padded(nic2)
        system_mac = manager_ports[0] if manager_ports else (system_port_macs[0] if system_port_macs else "")

        row = {
            "System Serial Number": self._clean_serial(system.get("SerialNumber") or chassis.get("SerialNumber") or ""),
            "Manufacturer": self._normalized(system.get("Manufacturer") or chassis.get("Manufacturer") or server.get("vendor") or ""),
            "Model": self._normalized(system.get("Model") or chassis.get("Model") or ""),
            "System MAC Address": system_mac,
            "IPMI IP": self._normalized(server.get("ip") or ""),
            "IPMI Username": self._normalized(server.get("ipmi_username") or ""),
            "IPMI Password": self._normalized(server.get("ipmi_password") or ""),
            "NIC1 MAC1": nic1[0],
            "NIC1 MAC2": nic1[1],
            "NIC1 MAC3": nic1[2],
            "NIC1 MAC4": nic1[3],
            "NIC2 MAC1": nic2[0],
            "NIC2 MAC2": nic2[1],
            "NIC2 MAC3": nic2[2],
            "NIC2 MAC4": nic2[3],
        }
        return row

    def write_sample_server_xlsx(self, path_xlsx: Path, servers: list[dict]):
        headers = self._server_sample_headers()
        path_xlsx.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "sample_server"
        ws.append(headers)

        for server in servers:
            row = self._build_sample_server_row(server)
            ws.append([row.get(header, "") for header in headers])

        self._style_sheet(ws, header_color="9E2A2B", freeze="A2", table_name="SampleServer")
        wb.save(path_xlsx)

    def write_serials_excel(self, path_xlsx: Path, servers: list[dict]):
        wb = Workbook()

        ws_overview = wb.active
        ws_overview.title = "Server Overview"
        ws_overview.append([
            "IP", "System Serial", "Chassis / Platform", "Motherboard",
            "CPU Qty", "DIMM Qty", "Drive Qty", "NIC Qty", "PSU Qty",
            "BIOS", "BMC/iLO"
        ])

        ws_detail = wb.create_sheet("Component Serials")
        ws_detail.append([
            "IP", "System Serial", "Category", "Part Number", "Description", "Serial / MAC"
        ])

        for server in servers:
            ip = server.get("ip", "")
            system = server.get("system") or {}
            chassis = server.get("chassis") or {}
            sys_sn = self._clean_serial(system.get("SerialNumber") or chassis.get("SerialNumber"))
            chassis_platform = chassis.get("Model") or system.get("Model") or ""
            motherboard = self._clean_part_number(system.get("PartNumber") or chassis.get("PartNumber") or "")

            cpu_qty = len([p for p in (server.get("procs") or []) if self._normalized(p.get("Model"))])
            dimm_items = [d for d in (server.get("dimms") or []) if isinstance(d.get("CapacityMiB"), int) and d.get("CapacityMiB") > 0 and self._clean_part_number(d.get("PartNumber") or d.get("Manufacturer") or "")]
            drive_items = [d for d in (server.get("drives") or []) if self._normalized(d.get("Model") or d.get("PartNumber"))]
            nic_items = [n for n in (server.get("nic_adapters") or []) if self._normalized(n.get("Model") or n.get("Name") or n.get("PartNumber"))]
            psu_items = [p for p in (server.get("psus") or []) if self._normalized(p.get("Model") or p.get("CapacityWatts"))]

            ws_overview.append([
                ip, sys_sn, chassis_platform, motherboard,
                cpu_qty, len(dimm_items), len(drive_items), len(nic_items), len(psu_items),
                self._normalized(server.get("bios")), self._normalized(server.get("bmcver")),
            ])

            def add_serial_row(category: str, pn: str, desc: str, serial: str):
                serial = self._clean_serial(serial)
                if not serial:
                    return
                ws_detail.append([ip, sys_sn, category, self._clean_part_number(pn), self._normalized(desc), serial])

            for proc in server.get("procs") or []:
                add_serial_row("CPU", proc.get("PartNumber") or "", proc.get("Model") or "CPU", proc.get("SerialNumber"))
            for dimm in dimm_items:
                desc = f"{int((dimm.get('CapacityMiB') or 0)/1024)}GB DIMM"
                add_serial_row("Memory", dimm.get("PartNumber") or dimm.get("Manufacturer") or "", desc, dimm.get("SerialNumber"))
            for drive in drive_items:
                model = drive.get("Model") or "Drive"
                add_serial_row("Drive", drive.get("PartNumber") or "", model, drive.get("_BetterSerial") or drive.get("SerialNumber"))
            for nic in server.get("nic_adapters") or []:
                add_serial_row("NIC", nic.get("PartNumber") or "", nic.get("Model") or nic.get("Name") or "NIC", nic.get("SerialNumber"))
            for port in server.get("nic_ports") or []:
                add_serial_row("NIC MAC", "", port.get("Name") or "MAC", port.get("MACAddress"))
            for psu in server.get("psus") or []:
                desc = f"{psu.get('CapacityWatts') or ''}W {self._normalized(psu.get('Model'))}".strip()
                add_serial_row("PSU", psu.get("Model") or "", desc, psu.get("SerialNumber"))

        self._style_sheet(ws_overview, header_color="1F4E78", table_name="tblServerOverview")
        self._style_sheet(ws_detail, header_color="375623", table_name="tblComponentSerials")
        path_xlsx.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path_xlsx)

    def _append_part_rows(self, rows: list[dict], server: dict):
        ip = server.get("ip", "")
        system = server.get("system") or {}
        chassis = server.get("chassis") or {}
        chassis_platform = chassis.get("Model") or system.get("Model") or ""
        motherboard = self._clean_part_number(system.get("PartNumber") or chassis.get("PartNumber") or "")
        serial = self._clean_serial(system.get("SerialNumber") or chassis.get("SerialNumber") or "")

        def add_row(category: str, part_number: str, description: str, qty: int):
            part_number = self._clean_part_number(part_number)
            description = self._normalized(description)
            if not part_number and not description:
                return
            rows.append({
                "IP": ip,
                "System Serial": serial,
                "Chassis / Platform": chassis_platform,
                "Motherboard": motherboard,
                "Category": category,
                "Part Number": part_number,
                "Description": description,
                "Quantity": qty,
            })

        if motherboard:
            add_row("Motherboard", motherboard, chassis_platform or motherboard, 1)

        cpu_counts = Counter()
        cpu_desc = {}
        for proc in server.get("procs") or []:
            pn = self._clean_part_number(proc.get("PartNumber") or "")
            model = self._normalized(proc.get("Model") or "CPU")
            key = pn or f"MODEL:{model}"
            cpu_counts[key] += 1
            cpu_desc[key] = model
        for key, qty in cpu_counts.items():
            add_row("CPU", "" if key.startswith("MODEL:") else key, cpu_desc.get(key, "CPU"), qty)

        mem_counts = Counter()
        mem_desc = {}
        for dimm in server.get("dimms") or []:
            cap = dimm.get("CapacityMiB") or 0
            pn = self._clean_part_number(dimm.get("PartNumber") or dimm.get("Manufacturer") or "")
            if not cap or not pn:
                continue
            key = pn
            mem_counts[key] += 1
            mem_desc[key] = f"{int(cap/1024)}GB DIMM"
        for key, qty in mem_counts.items():
            add_row("Memory", key, mem_desc.get(key, "DIMM"), qty)

        drive_counts = Counter()
        drive_desc = {}
        for drive in server.get("drives") or []:
            pn = self._clean_part_number(drive.get("PartNumber") or "")
            model = self._normalized(drive.get("Model") or "Drive")
            size = drive.get("CapacityBytes")
            size_text = f"{int(round(size/(1024**3)))}GB " if isinstance(size, int) and size > 0 else ""
            key = pn or f"MODEL:{model}"
            drive_counts[key] += 1
            drive_desc[key] = f"{size_text}{model}".strip()
        for key, qty in drive_counts.items():
            add_row("Drive", "" if key.startswith("MODEL:") else key, drive_desc.get(key, "Drive"), qty)

        nic_counts = Counter()
        nic_desc = {}
        for adapter in server.get("nic_adapters") or []:
            model = self._normalized(adapter.get("Model") or adapter.get("Name") or "NIC")
            pn = self._clean_part_number(adapter.get("PartNumber") or "")
            if not pn:
                oem = adapter.get("Oem") or {}
                for section in oem.values():
                    if isinstance(section, dict):
                        for name, value in section.items():
                            if isinstance(value, str) and ("part" in name.lower() or "spn" in name.lower()):
                                pn = self._clean_part_number(value)
                                break
                    if pn:
                        break
            key = pn or f"MODEL:{model}"
            nic_counts[key] += 1
            nic_desc[key] = model
        for key, qty in nic_counts.items():
            add_row("NIC", "" if key.startswith("MODEL:") else key, nic_desc.get(key, "NIC"), qty)

        psu_counts = Counter()
        psu_desc = {}
        for psu in server.get("psus") or []:
            model = self._clean_part_number(psu.get("Model") or "")
            watts = psu.get("CapacityWatts")
            desc = f"{watts}W {model}".strip() if watts or model else ""
            key = model or f"MODEL:{desc or 'PSU'}"
            if not key:
                continue
            psu_counts[key] += 1
            psu_desc[key] = desc or model or "PSU"
        for key, qty in psu_counts.items():
            add_row("PSU", "" if key.startswith("MODEL:") else key, psu_desc.get(key, "PSU"), qty)


    def write_part_inventory_excel(self, path_xlsx: Path, servers: list[dict]):
        detail_rows: list[dict] = []
        for server in servers:
            self._append_part_rows(detail_rows, server)

        grand = defaultdict(lambda: {"Quantity": 0, "Platforms": set(), "Boards": set(), "ServerCount": set()})
        by_config = defaultdict(lambda: {"Quantity": 0, "ServerCount": set()})
        category_totals = defaultdict(lambda: {"Quantity": 0, "PartNumbers": set()})

        for row in detail_rows:
            grand_key = (row["Category"], row["Part Number"], row["Description"])
            grand[grand_key]["Quantity"] += int(row["Quantity"] or 0)
            if row["Chassis / Platform"]:
                grand[grand_key]["Platforms"].add(row["Chassis / Platform"])
            if row["Motherboard"]:
                grand[grand_key]["Boards"].add(row["Motherboard"])
            if row["IP"]:
                grand[grand_key]["ServerCount"].add(row["IP"])

            cfg_key = (row["Chassis / Platform"], row["Motherboard"], row["Category"], row["Part Number"], row["Description"])
            by_config[cfg_key]["Quantity"] += int(row["Quantity"] or 0)
            if row["IP"]:
                by_config[cfg_key]["ServerCount"].add(row["IP"])

            category_totals[row["Category"]]["Quantity"] += int(row["Quantity"] or 0)
            if row["Part Number"]:
                category_totals[row["Category"]]["PartNumbers"].add(row["Part Number"])

        variant_groups = self.hardware_groups(servers)

        wb = Workbook()

        # Dashboard
        ws_dash = wb.active
        ws_dash.title = "Dashboard"
        ws_dash["A1"] = "Server Intake Inventory Summary"
        ws_dash["A1"].font = Font(size=16, bold=True)
        ws_dash["A3"] = "Generated"
        ws_dash["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        metrics = [
            ("Total Servers", len(servers)),
            ("Config Variants", len(variant_groups)),
            ("Distinct Part Numbers", len({k[1] for k in grand.keys() if k[1]})),
            ("Total CPU Qty", category_totals.get("CPU", {}).get("Quantity", 0)),
            ("Total Memory Qty", category_totals.get("Memory", {}).get("Quantity", 0)),
            ("Total Drive Qty", category_totals.get("Drive", {}).get("Quantity", 0)),
            ("Total NIC Qty", category_totals.get("NIC", {}).get("Quantity", 0)),
            ("Total PSU Qty", category_totals.get("PSU", {}).get("Quantity", 0)),
        ]
        row_idx = 5
        for label, value in metrics:
            ws_dash[f"A{row_idx}"] = label
            ws_dash[f"B{row_idx}"] = value
            ws_dash[f"A{row_idx}"].font = Font(bold=True)
            row_idx += 1
        row_idx += 1
        ws_dash[f"A{row_idx}"] = "Category"
        ws_dash[f"B{row_idx}"] = "Distinct PN"
        ws_dash[f"C{row_idx}"] = "Total Qty"
        for c in ws_dash[row_idx]:
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.font = Font(color="FFFFFF", bold=True)
        for cat in sorted(category_totals.keys()):
            row_idx += 1
            ws_dash[f"A{row_idx}"] = cat
            ws_dash[f"B{row_idx}"] = len(category_totals[cat]["PartNumbers"])
            ws_dash[f"C{row_idx}"] = category_totals[cat]["Quantity"]
        ws_dash.column_dimensions["A"].width = 28
        ws_dash.column_dimensions["B"].width = 16
        ws_dash.column_dimensions["C"].width = 14

        # Totals by Part
        ws_main = wb.create_sheet("Totals by Part")
        ws_main.append(["Category", "Part Number", "Description", "Total Qty", "Server Count", "Chassis / Platform(s)", "Motherboard(s)"])
        sorted_keys = sorted(grand.keys(), key=lambda k: (str(k[0]), -grand[k]["Quantity"], str(k[1]), str(k[2])))
        for key in sorted_keys:
            data = grand[key]
            ws_main.append([
                key[0], key[1], key[2], data["Quantity"], len(data["ServerCount"]),
                ", ".join(sorted(data["Platforms"])),
                ", ".join(sorted(data["Boards"])),
            ])
        self._style_sheet(ws_main, header_color="1F4E78", table_name="tblTotalsByPart")

        # Category Summary
        ws_cat = wb.create_sheet("Category Summary")
        ws_cat.append(["Category", "Distinct Part Numbers", "Total Qty"])
        for category in sorted(category_totals.keys()):
            ws_cat.append([category, len(category_totals[category]["PartNumbers"]), category_totals[category]["Quantity"]])
        self._style_sheet(ws_cat, header_color="7F6000", table_name="tblCategorySummary")

        # Config Summary
        ws_cfgsum = wb.create_sheet("Config Summary")
        ws_cfgsum.append(["Config", "Server Count", "IPs", "Serials", "Platform", "Motherboard", "CPU", "Memory", "Drives", "NIC", "PSU", "BIOS", "BMC/iLO"])
        for idx, group in enumerate(variant_groups, start=1):
            ref = group[0]
            system = ref.get("system") or {}
            chassis = ref.get("chassis") or {}
            ws_cfgsum.append([
                f"Config {idx}",
                len(group),
                ", ".join(s.get("ip", "") for s in group if s.get("ip")),
                ", ".join(self._clean_serial((s.get("system") or {}).get("SerialNumber") or (s.get("chassis") or {}).get("SerialNumber")) for s in group if self._clean_serial((s.get("system") or {}).get("SerialNumber") or (s.get("chassis") or {}).get("SerialNumber"))),
                chassis.get("Model") or system.get("Model") or "",
                self._clean_part_number(system.get("PartNumber") or chassis.get("PartNumber") or ""),
                self.summarize_cpu(ref.get("procs") or []),
                self.summarize_memory(ref.get("dimms") or []),
                self.summarize_drives(ref.get("drives") or []),
                self.summarize_network(ref.get("nic_adapters") or []),
                self.summarize_psus(ref.get("psus") or []),
                self._variant_summary([self._normalized(s.get("bios")) for s in group]),
                self._variant_summary([self._normalized(s.get("bmcver")) for s in group]),
            ])
        self._style_sheet(ws_cfgsum, header_color="375623", table_name="tblConfigSummary")

        # By Platform+Board
        ws_cfg = wb.create_sheet("By Platform+Board")
        ws_cfg.append(["Chassis / Platform", "Motherboard", "Category", "Part Number", "Description", "Total Qty", "Server Count"])
        for key in sorted(by_config.keys(), key=lambda k: (str(k[0]), str(k[1]), str(k[2]), -by_config[k]["Quantity"], str(k[3]))):
            data = by_config[key]
            ws_cfg.append([key[0], key[1], key[2], key[3], key[4], data["Quantity"], len(data["ServerCount"])])
        self._style_sheet(ws_cfg, header_color="5B9BD5", table_name="tblByPlatformBoard")

        # Detail
        ws_detail = wb.create_sheet("Per Server Detail")
        ws_detail.append(["IP", "System Serial", "Chassis / Platform", "Motherboard", "Category", "Part Number", "Description", "Quantity"])
        for row in sorted(detail_rows, key=lambda r: (r["IP"], r["Category"], str(r["Part Number"]), str(r["Description"]))):
            ws_detail.append([row["IP"], row["System Serial"], row["Chassis / Platform"], row["Motherboard"], row["Category"], row["Part Number"], row["Description"], row["Quantity"]])
        self._style_sheet(ws_detail, header_color="A61C00", table_name="tblPerServerDetail")

        path_xlsx.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path_xlsx)

    def _counts_cpu(self, server: dict) -> Counter:
        return Counter([p.get("Model") for p in server.get("procs") or [] if p.get("Model")])

    def _counts_memory(self, server: dict) -> Counter:
        counts = Counter()
        for dimm in server.get("dimms") or []:
            cap = dimm.get("CapacityMiB")
            pn = self._clean_part_number(dimm.get("PartNumber") or dimm.get("Manufacturer") or "")
            if not isinstance(cap, int) or cap <= 0 or not pn:
                continue
            size_gb = f"{int(cap/1024)}GB"
            key = f"{size_gb} {pn}".strip()
            counts[key] += 1
        return counts

    def _counts_drives(self, server: dict) -> Counter:
        counts = Counter()
        for drive in server.get("drives") or []:
            model = self._normalized(drive.get("Model") or "Drive")
            size = drive.get("CapacityBytes")
            size_gb = ""
            if isinstance(size, int) and size > 0:
                size_gb = f"{int(round(size/(1024**3)))}GB"
            pn = drive.get("PartNumber") or ""
            key = f"{size_gb} {model}".strip()
            if pn:
                key += f" ({pn})"
            counts[key] += 1
        return counts

    def _counts_nic(self, server: dict) -> Counter:
        counts = Counter()
        for adapter in server.get("nic_adapters") or []:
            model = adapter.get("Model") or adapter.get("Name") or "NIC"
            pn = adapter.get("PartNumber") or ""
            if not pn:
                oem = adapter.get("Oem") or {}
                for section in oem.values():
                    if isinstance(section, dict):
                        for name, value in section.items():
                            if isinstance(value, str) and ("part" in name.lower() or "spn" in name.lower()):
                                pn = value
                                break
                    if pn:
                        break
            key = f"{model} ({pn})" if pn else model
            counts[key] += 1
        return counts

    def _counts_psu(self, server: dict) -> Counter:
        counts = Counter()
        for psu in server.get("psus") or []:
            watts = psu.get("CapacityWatts")
            model = self._clean_part_number(psu.get("Model") or "")
            if not watts and not model:
                continue
            key = f"{watts}W {model}".strip() if watts else model
            counts[key] += 1
        return counts

    def fingerprint_for_group(self, server: dict):
        system = server.get("system") or {}
        chassis = server.get("chassis") or {}
        mb_pn = system.get("PartNumber") or chassis.get("PartNumber") or ""
        chassis_model = chassis.get("Model") or system.get("Model") or ""
        return (
            tuple(sorted(self._counts_cpu(server).items())),
            tuple(sorted(self._counts_memory(server).items())),
            tuple(sorted(self._counts_drives(server).items())),
            tuple(sorted(self._counts_nic(server).items())),
            tuple(sorted(self._counts_psu(server).items())),
            mb_pn,
            chassis_model,
        )

    def configs_all_equal(self, servers: list[dict]) -> bool:
        if not servers:
            return True
        baseline = self.fingerprint_for_group(servers[0])
        return all(self.fingerprint_for_group(server) == baseline for server in servers[1:])

    def hardware_groups(self, servers: list[dict]):
        grouped = defaultdict(list)
        for server in servers:
            grouped[self.fingerprint_for_group(server)].append(server)
        groups = sorted(grouped.values(), key=lambda items: (-len(items), items[0].get("ip", "")))
        return groups

    def choose_majority_server(self, servers: list[dict]) -> tuple[dict, int, bool]:
        groups = self.hardware_groups(servers)
        top_count = len(groups[0])
        next_count = len(groups[1]) if len(groups) > 1 else 0
        has_majority = top_count > next_count
        return groups[0][0], top_count, has_majority

    def _format_counter(self, counter: Counter) -> str:
        if not counter:
            return "none"
        parts = []
        for key, qty in sorted(counter.items(), key=lambda item: str(item[0])):
            parts.append(f"{qty}x {key}")
        return ", ".join(parts)

    def _variant_summary(self, values: list[str]) -> str:
        cleaned = [self._normalized(v) for v in values if self._normalized(v)]
        unique = sorted(set(cleaned))
        if not unique:
            return ""
        if len(unique) == 1:
            return unique[0]
        counts = Counter(cleaned)
        return " | ".join(f"{qty} srv: {val}" for val, qty in sorted(counts.items(), key=lambda item: (-item[1], item[0])))

    def _config_label(self, idx: int) -> str:
        return f"Config {idx}"


    def _variant_dict(self, group: list[dict], idx: int) -> dict:
        ref = group[0]
        system = ref.get("system") or {}
        chassis = ref.get("chassis") or {}
        platform = chassis.get("Model") or system.get("Model") or ""
        motherboard = self._clean_part_number(system.get("PartNumber") or chassis.get("PartNumber") or "")
        return {
            "name": f"Config {idx}",
            "count": len(group),
            "ips": [s.get("ip", "") for s in group if s.get("ip")],
            "serials": [self._clean_serial((s.get("system") or {}).get("SerialNumber") or (s.get("chassis") or {}).get("SerialNumber")) for s in group if self._clean_serial((s.get("system") or {}).get("SerialNumber") or (s.get("chassis") or {}).get("SerialNumber"))],
            "platform": platform,
            "motherboard": motherboard,
            "cpu": self.summarize_cpu(ref.get("procs") or []) or "none",
            "memory": self.summarize_memory(ref.get("dimms") or []) or "none",
            "drives": self.summarize_drives(ref.get("drives") or []) or "none",
            "nic": self.summarize_network(ref.get("nic_adapters") or []) or "none",
            "psu": self.summarize_psus(ref.get("psus") or []) or "none",
            "bios": self._variant_summary([self._normalized(s.get("bios")) for s in group]) or "none",
            "bmc": self._variant_summary([self._normalized(s.get("bmcver")) for s in group]) or "none",
        }

    def build_group_report(self, servers: list[dict]) -> dict:
        groups = self.hardware_groups(servers)
        variants = [self._variant_dict(group, idx) for idx, group in enumerate(groups, start=1)]
        top_count = variants[0]["count"] if variants else 0
        next_count = variants[1]["count"] if len(variants) > 1 else 0
        has_majority = top_count > next_count
        if len(variants) == 1:
            headline = f"Single hardware configuration across {len(servers)} server{'s' if len(servers) != 1 else ''}"
        elif has_majority:
            headline = f"Primary configuration detected: {top_count}/{len(servers)} servers, {len(variants)} total variants"
        else:
            headline = f"Multiple hardware configurations detected: {len(variants)} variants across {len(servers)} servers"

        categories = Counter()
        for server in servers:
            categories["CPU"] += sum(self._counts_cpu(server).values())
            categories["Memory"] += sum(self._counts_memory(server).values())
            categories["Drive"] += sum(self._counts_drives(server).values())
            categories["NIC"] += sum(self._counts_nic(server).values())
            categories["PSU"] += sum(self._counts_psu(server).values())

        return {
            "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "server_count": len(servers),
            "variant_count": len(variants),
            "headline": headline,
            "has_majority": has_majority,
            "variants": variants,
            "category_totals": dict(categories),
        }

    def _draw_group_report_page(self, c, report: dict, page_width: float, page_height: float):
        margin = 34
        y = page_height - 36

        def ensure_space(needed=18):
            nonlocal y
            if y < margin + needed:
                c.showPage()
                y = page_height - 36
                c.setFont("Helvetica", 10)

        c.setFont("Helvetica-Bold", 16)
        c.drawString(margin, y, "Configurator Group Summary")
        y -= 18
        c.setFont("Helvetica", 9)
        c.drawString(margin, y, f"Generated: {report['generated_at']}")
        y -= 18
        c.setFont("Helvetica-Bold", 11)
        c.drawString(margin, y, report["headline"])
        y -= 22

        # metrics
        c.setFont("Helvetica-Bold", 10)
        c.drawString(margin, y, "Batch Totals")
        y -= 16
        c.setFont("Helvetica", 9)
        metrics = [
            f"Servers: {report['server_count']}",
            f"Variants: {report['variant_count']}",
            f"CPU: {report['category_totals'].get('CPU',0)}",
            f"Memory: {report['category_totals'].get('Memory',0)}",
            f"Drives: {report['category_totals'].get('Drive',0)}",
            f"NIC: {report['category_totals'].get('NIC',0)}",
            f"PSU: {report['category_totals'].get('PSU',0)}",
        ]
        c.drawString(margin, y, " | ".join(metrics))
        y -= 20

        # summary table-ish list
        c.setFont("Helvetica-Bold", 10)
        c.drawString(margin, y, "Configuration Overview")
        y -= 16
        for variant in report["variants"]:
            ensure_space(42)
            c.setFont("Helvetica-Bold", 9)
            c.drawString(margin, y, f"{variant['name']}  |  {variant['count']} server(s)  |  {variant['platform'] or 'n/a'}")
            y -= 12
            c.setFont("Helvetica", 8.5)
            summary = f"Board: {variant['motherboard'] or 'n/a'} | CPU: {variant['cpu']} | Memory: {variant['memory']}"
            for line in simpleSplit(summary, "Helvetica", 8.5, page_width - 2 * margin):
                c.drawString(margin + 10, y, line)
                y -= 11
            y -= 4

        y -= 8
        for variant in report["variants"]:
            ensure_space(120)
            c.setFont("Helvetica-Bold", 11)
            c.drawString(margin, y, f"{variant['name']} Details")
            y -= 14
            c.setFont("Helvetica", 9)
            details = [
                f"IPs: {', '.join(variant['ips']) or 'n/a'}",
                f"Serials: {', '.join(variant['serials']) or 'n/a'}",
                f"Platform: {variant['platform'] or 'n/a'}",
                f"Motherboard: {variant['motherboard'] or 'n/a'}",
                f"CPU: {variant['cpu']}",
                f"Memory: {variant['memory']}",
                f"Drives: {variant['drives']}",
                f"NIC: {variant['nic']}",
                f"PSU: {variant['psu']}",
                f"BIOS: {variant['bios']}",
                f"BMC/iLO: {variant['bmc']}",
            ]
            for item in details:
                ensure_space(14)
                for line in simpleSplit(item, "Helvetica", 9, page_width - 2 * margin - 12):
                    c.drawString(margin + 8, y, line)
                    y -= 11
            y -= 8

    def write_group_summary_pdf(self, report: dict, output_pdf: Path):
        output_pdf.parent.mkdir(parents=True, exist_ok=True)
        page_width = 792
        page_height = 612
        c = canvas.Canvas(str(output_pdf), pagesize=(page_width, page_height))
        self._draw_group_report_page(c, report, page_width, page_height)
        c.save()

    def write_group_summary_editable_pdf(self, report: dict, output_pdf: Path):
        output_pdf.parent.mkdir(parents=True, exist_ok=True)
        page_width = 792
        page_height = 612
        c = canvas.Canvas(str(output_pdf), pagesize=(page_width, page_height))
        form = c.acroForm
        margin = 34
        y = page_height - 36

        def label(x, y, text, size=10):
            c.setFont("Helvetica-Bold", size)
            c.drawString(x, y, text)

        def field(name, x, y, w, h, value="", size=9):
            text_value = str(value or "")
            c.setFont("Helvetica", size)
            lines = simpleSplit(text_value, "Helvetica", size, max(w - 8, 20))
            baseline = y + h - size - 4
            for line in lines[: max(int((h - 8) / (size + 2)), 1)]:
                c.drawString(x + 4, baseline, line)
                baseline -= size + 2
            form.textfield(name=name, x=x, y=y, width=w, height=h, value=text_value, fontName="Helvetica", fontSize=size, borderStyle="solid", borderWidth=0.6, forceBorder=True, textColor=colors.black, borderColor=colors.HexColor("#777777"), fillColor=None, fieldFlags=4096)

        label(margin, y, "Configurator Group Summary (Editable)", 15)
        y -= 20
        c.setFont("Helvetica", 9)
        c.drawString(margin, y, f"Generated: {report['generated_at']}")
        y -= 22

        label(margin, y, "Overview", 10)
        overview_text = report['headline'] + "\n" + " | ".join([
            f"Servers: {report['server_count']}",
            f"Variants: {report['variant_count']}",
            f"CPU: {report['category_totals'].get('CPU',0)}",
            f"Memory: {report['category_totals'].get('Memory',0)}",
            f"Drives: {report['category_totals'].get('Drive',0)}",
            f"NIC: {report['category_totals'].get('NIC',0)}",
            f"PSU: {report['category_totals'].get('PSU',0)}",
        ])
        field("overview", margin, y - 78, page_width - 2*margin, 68, overview_text, 9)
        y -= 96

        for idx, variant in enumerate(report["variants"], start=1):
            if y < 180:
                c.showPage()
                form = c.acroForm
                y = page_height - 36
            label(margin, y, f"{variant['name']} ({variant['count']} server(s))", 10)
            block = "\n".join([
                f"IPs: {', '.join(variant['ips']) or 'n/a'}",
                f"Serials: {', '.join(variant['serials']) or 'n/a'}",
                f"Platform: {variant['platform'] or 'n/a'}",
                f"Motherboard: {variant['motherboard'] or 'n/a'}",
                f"CPU: {variant['cpu']}",
                f"Memory: {variant['memory']}",
                f"Drives: {variant['drives']}",
                f"NIC: {variant['nic']}",
                f"PSU: {variant['psu']}",
                f"BIOS: {variant['bios']}",
                f"BMC/iLO: {variant['bmc']}",
            ])
            field(f"variant_{idx}", margin, y - 132, page_width - 2*margin, 120, block, 9)
            y -= 150

        c.save()
        self._finalize_fillable_pdf(output_pdf)

    def _group_detail_lines(self, servers: list[dict]) -> tuple[list[str], list[str], bool]:
        report = self.build_group_report(servers)
        overview = [report['headline']]
        details = []
        for variant in report['variants']:
            overview.append(f"{variant['name']}: {variant['count']} srv | {variant['platform'] or 'n/a'} | {variant['memory']}")
            details.extend([
                f"{variant['name']} ({variant['count']} server(s))",
                f"IPs: {', '.join(variant['ips']) or 'n/a'}",
                f"Serials: {', '.join(variant['serials']) or 'n/a'}",
                f"Platform: {variant['platform'] or 'n/a'}",
                f"Motherboard: {variant['motherboard'] or 'n/a'}",
                f"CPU: {variant['cpu']}",
                f"Memory: {variant['memory']}",
                f"Drives: {variant['drives']}",
                f"NIC: {variant['nic']}",
                f"PSU: {variant['psu']}",
                f"BIOS: {variant['bios']}",
                f"BMC/iLO: {variant['bmc']}",
                "",
            ])
        return overview, details, report['has_majority']

    def group_extras_differences(self, servers: list[dict], base: dict | None = None) -> tuple[str, str, bool]:
        if not servers:
            return "", "", True
        report = self.build_group_report(servers)
        overview_lines, detail_lines, _ = self._group_detail_lines(servers)
        return "\n".join(overview_lines).strip(), "\n".join(detail_lines).strip(), report['has_majority']

    def build_group_values(self, servers: list[dict]) -> tuple[dict, bool]:
        report = self.build_group_report(servers)
        base_server = self.choose_majority_server(servers)[0]
        values = self.describe_for_pdf(base_server)
        values["QUANTITY"] = str(len(servers))
        serials = [self._clean_serial((s.get("system") or {}).get("SerialNumber") or (s.get("chassis") or {}).get("SerialNumber")) for s in servers]
        values["SYSTEM_SERIALS"] = "; ".join([s for s in serials if s][:8]) if len([s for s in serials if s]) <= 8 else f"{len(servers)} servers (see serials.xlsx)"
        overview_lines, detail_lines, has_majority = self._group_detail_lines(servers)
        values["EXTRAS"] = "\n".join(overview_lines)
        values["NOTES"] = "\n".join(detail_lines)
        values["BIOS"] = self._variant_summary([self._normalized(s.get("bios")) for s in servers])
        values["BMC_ILO"] = self._variant_summary([self._normalized(s.get("bmcver")) for s in servers])
        values["MEMORY"] = self._variant_summary([self.summarize_memory(s.get("dimms") or []) for s in servers])
        values["DRIVES"] = self._variant_summary([self.summarize_drives(s.get("drives") or []) for s in servers]) or "none"
        values["NETWORK_HW"] = self._variant_summary([self.summarize_network(s.get("nic_adapters") or []) for s in servers]) or "none"
        values["PSU"] = self._variant_summary([self.summarize_psus(s.get("psus") or []) for s in servers])
        values["CHASSIS_PLATFORM"] = self._variant_summary([(s.get("chassis") or {}).get("Model") or (s.get("system") or {}).get("Model") or "" for s in servers])
        values["MOTHERBOARD"] = self._variant_summary([self._clean_part_number((s.get("system") or {}).get("PartNumber") or (s.get("chassis") or {}).get("PartNumber") or "") for s in servers])
        values["CPU"] = self._variant_summary([self.summarize_cpu(s.get("procs") or []) for s in servers])
        return values, has_majority

    def collect_one(self, row: TaskRow, vendor_key: str, options: RunOptions, log: LogFn) -> dict | None:
        ip = row.ip
        auth = (row.user, row.password)
        root = self.rf_get(ip, "/redfish/v1", auth, min(10, options.timeout_seconds), log)
        if not root:
            log(f"[{ip}] Redfish root unavailable")
            return None

        sys_path, ch_path, mgr_path = self.discover_paths(ip, auth, vendor_key, options.timeout_seconds, log)
        system = self.get_system(ip, auth, sys_path, options.timeout_seconds, log)
        chassis = self.get_chassis(ip, auth, ch_path, options.timeout_seconds, log)
        return {
            "ip": ip,
            "vendor": (system.get("Manufacturer") or vendor_key or "").lower(),
            "ipmi_username": row.user,
            "ipmi_password": row.password,
            "system": system,
            "chassis": chassis,
            "procs": self.get_processors(ip, auth, sys_path, options.timeout_seconds, log),
            "dimms": self.get_dimms(ip, auth, sys_path, options.timeout_seconds, log),
            "drives": self.get_drives(ip, auth, sys_path, ch_path, options.timeout_seconds, log),
            "nic_ports": self.get_nic_ports(ip, auth, sys_path, mgr_path, options.timeout_seconds, log),
            "manager_ports": self.get_manager_ports(ip, auth, mgr_path, options.timeout_seconds, log),
            "nic_adapters": self.get_nic_adapters(ip, auth, sys_path, ch_path, options.timeout_seconds, log),
            "psus": self.get_psus(ip, auth, ch_path, options.timeout_seconds, log),
            "bios": self.get_bios_version(system),
            "bmcver": self.get_bmc_version(ip, auth, mgr_path, options.timeout_seconds, log),
        }

    def output_root_for_vendor(self, vendor_key: str) -> Path:
        mapping = {
            "dell": self.base_dir / "Configurator_PDFs_Dell",
            "hpe": self.base_dir / "Configurator_PDFs",
            "supermicro": self.base_dir / "Configurator_PDFs_Supermicro",
        }
        root = mapping.get(vendor_key, mapping["hpe"])
        root.mkdir(parents=True, exist_ok=True)
        return root

    def run(self, tasks: Iterable[TaskRow], vendor_key: str, folder_name: str, output_root: Path | None,
            options: RunOptions, log: LogFn, progress: ProgressFn, status: StatusFn) -> Path:
        rows = [row for row in tasks if row.enabled and row.ip.strip()]
        if not rows:
            raise BackendError("No enabled task rows to run.")
        if not folder_name.strip():
            folder_name = datetime.now().strftime("RUN_%Y%m%d_%H%M%S")

        out_root = Path(output_root) if output_root else self.output_root_for_vendor(vendor_key)
        out_dir = out_root / folder_name.strip()
        out_dir.mkdir(parents=True, exist_ok=True)

        servers: list[dict] = []
        total = len(rows)
        for index, row in enumerate(rows, start=1):
            progress(index - 1, total)
            status(row.ip, "Running")
            log(f"[{row.ip}] Collecting Redfish data...")
            try:
                data = self.collect_one(row, vendor_key, options, log)
                if not data:
                    status(row.ip, "Unreachable")
                    if not options.skip_unreachable:
                        raise BackendError(f"{row.ip} unreachable")
                    continue
                servers.append(data)
                status(row.ip, "Success")
                if options.create_individual_pdfs:
                    values = self.describe_for_pdf(data)
                    pdf_path = out_dir / f"Configurator_{row.ip.replace(':', '_')}.pdf"
                    self.write_pdf_from_values(values, pdf_path)
            except Exception as exc:
                status(row.ip, f"Error: {exc}")
                log(f"[{row.ip}] ERROR: {exc}")

        progress(total, total)
        if not servers:
            raise BackendError("No server data collected.")

        if options.create_group_pdf:
            report = self.build_group_report(servers)
            self.write_group_summary_pdf(report, out_dir / "Configurator_GROUP.pdf")
            self.write_group_summary_editable_pdf(report, out_dir / "Configurator_GROUP_editable.pdf")

        if options.create_serial_excel:
            self.write_serials_excel(out_dir / f"{folder_name}_serials.xlsx", servers)
            self.write_part_inventory_excel(out_dir / f"{folder_name}_part_numbers.xlsx", servers)

        self.write_sample_server_xlsx(out_dir / f"{folder_name}_sample_server.xlsx", servers)

        log(f"Output ready: {out_dir}")
        return out_dir
